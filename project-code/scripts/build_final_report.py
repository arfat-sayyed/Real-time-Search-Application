from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs' / 'final-project-report.md'
OUTPUT = ROOT / 'Final Project Report Findly.docx'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), fill)
    tc_pr.append(shading)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader')
    repeat.set(qn('w:val'), 'true')
    tr_pr.append(repeat)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    run.font.size = Pt(10)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run('Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(field)


def add_toc(paragraph):
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
    paragraph._p.append(field)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    for style in ['Heading 1', 'Heading 2', 'Heading 3']:
        s = doc.styles[style]
        s.font.name = 'Times New Roman'
        s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        s.font.color.rgb = RGBColor(0, 0, 0)


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run('FINDLY REAL TIME PRODUCT SEARCH')
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.bold = True
    add_page_field(section.footer.paragraphs[0])


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run('AMITY UNIVERSITY ONLINE\nNOIDA, UTTAR PRADESH')
    r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(14)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(32)
    r = p.add_run('DEVELOPING A REAL TIME PRODUCT SEARCH APPLICATION\nWITH NODE JS REACT AND ELASTICSEARCH')
    r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(16)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(32)
    p.add_run('A Project Report submitted in partial fulfillment of the requirements for the award of the degree of Master of Business Administration').font.size = Pt(12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(35)
    r = p.add_run('Submitted by\n[STUDENT NAME]\nEnrollment No. [ENROLLMENT NUMBER]\n\nGuided by\n[MENTOR NAME]\n\nAcademic Year [YEAR]')
    r.font.size = Pt(12)
    doc.add_page_break()


def add_declaration(doc):
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run('DECLARATION'); r.bold = True; r.font.size = Pt(14)
    p = doc.add_paragraph()
    p.add_run('I, [STUDENT NAME], a student pursuing [PROGRAM AND SEMESTER] at Amity University Online, hereby declare that the project work entitled “Developing a Real Time Product Search Application with Node JS React and Elasticsearch” has been prepared by me during the academic year [YEAR] under the guidance of [MENTOR NAME]. I certify that this is original bona fide work carried out by me and that it has not been submitted to any other university for the award of any degree.').font.size = Pt(12)
    doc.add_paragraph('\n\nSignature of Student: ______________________________')
    doc.add_page_break()


def add_front_matter(doc):
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run('TABLE OF CONTENTS'); r.bold = True; r.font.size = Pt(14)
    for item in [
        'Declaration', 'List of Tables', 'List of Figures', 'Chapter 1 Introduction to the Topic',
        'Chapter 2 Review of Literature', 'Chapter 3 Research Objectives and Methodology',
        'Chapter 4 Data Analysis and Results', 'Chapter 5 Findings and Conclusion',
        'Chapter 6 Recommendations and Limitations of the Study', 'Chapter 7 Bibliography and References',
        'Appendix A System Design and API', 'Appendix B Testing Evidence', 'Screenshot Appendix'
    ]:
        doc.add_paragraph(item)
    doc.add_paragraph('Update page numbers in Word after you finish adding screenshots.')
    doc.add_page_break()
    for title, items in [
        ('LIST OF TABLES', ['Table 1. Technology selection', 'Table 2. Functional test scenarios and results']),
        ('LIST OF FIGURES', ['Figure 1. Application architecture', 'Figure 2. Main product search interface', 'Figure 3. Autocomplete suggestions', 'Figure 4. Filtered search results'])
    ]:
        h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = h.add_run(title); r.bold = True; r.font.size = Pt(14)
        for item in items: doc.add_paragraph(item)
        doc.add_page_break()


def add_table(doc, rows):
    values = [[x.strip() for x in row.strip().strip('|').split('|')] for row in rows]
    values = [row for row in values if row and not all(set(cell) <= {'-', ':'} for cell in row)]
    table = doc.add_table(rows=0, cols=len(values[0]))
    table.style = 'Table Grid'
    for index, row_values in enumerate(values):
        cells = table.add_row().cells
        for col, value in enumerate(row_values):
            set_cell_text(cells[col], value.replace('`', ''), bold=index == 0, color=(255,255,255) if index == 0 else None)
            if index == 0: set_cell_shading(cells[col], '1F4E79')
        if index == 0: set_repeat_table_header(table.rows[0])
    doc.add_paragraph()


def add_markdown_content(doc):
    lines = SOURCE.read_text().splitlines()
    i = 0
    in_body = False
    while i < len(lines):
        line = lines[i].strip()
        if line == '# Chapter 1 Introduction to the Topic':
            in_body = True
        if not in_body:
            i += 1
            continue
        if not line or line == '---':
            i += 1; continue
        if line.startswith('# '):
            text = line[2:]
            if text.startswith('Developing a Real'):
                i += 1; continue
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(14)
            i += 1; continue
        if line.startswith('## '):
            text = line[3:]
            p = doc.add_paragraph()
            if text.startswith('Chapter '):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(14)
            else:
                r = p.add_run(text); r.bold = True; r.font.size = Pt(12)
            i += 1; continue
        if line.startswith('### '):
            p = doc.add_paragraph(); r = p.add_run(line[4:]); r.bold = True; r.font.size = Pt(12)
            i += 1; continue
        if line.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i]); i += 1
            add_table(doc, table_lines); continue
        if line.startswith('```'):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i]); i += 1
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run('\n'.join(code)); run.font.name = 'Courier New'; run.font.size = Pt(10)
            i += 1; continue
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet'); p.add_run(line[2:]); i += 1; continue
        if line[:2].isdigit() and '. ' in line[:5]:
            p = doc.add_paragraph(style='List Number'); p.add_run(line.split('. ', 1)[1]); i += 1; continue
        p = doc.add_paragraph()
        p.add_run(line.replace('**', '').replace('*', '').replace('`', ''))
        i += 1


def add_screenshot_slots(doc):
    doc.add_page_break()
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run('SCREENSHOT APPENDIX'); r.bold = True; r.font.size = Pt(14)
    doc.add_paragraph('Paste your own screenshots in the marked areas below. Keep the caption directly below each image and label the source as “Author’s screenshot.”')
    slots = [
        'Figure A1. Main Findly home screen showing 100 products and Indian rupee prices.',
        'Figure A2. Autocomplete suggestion dropdown while searching a product.',
        'Figure A3. Filtered search results using category and price range.',
        'Figure A4. Docker Desktop or terminal evidence of the Elasticsearch container.',
        'Figure A5. API health response and successful automated test output.'
    ]
    for label in slots:
        box = doc.add_table(rows=1, cols=1)
        box.style = 'Table Grid'
        cell = box.cell(0, 0)
        set_cell_shading(cell, 'F2F2F2')
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(36); p.paragraph_format.space_after = Pt(36)
        r = p.add_run('[PASTE SCREENSHOT HERE]'); r.italic = True; r.font.size = Pt(11)
        caption = doc.add_paragraph(label); caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        source = doc.add_paragraph('Source: Author’s screenshot.'); source.alignment = WD_ALIGN_PARAGRAPH.CENTER


doc = Document()
style_document(doc)
add_header_footer(doc.sections[0])
add_cover(doc)
add_declaration(doc)
add_front_matter(doc)
add_markdown_content(doc)
add_screenshot_slots(doc)
doc.save(OUTPUT)
print(OUTPUT)
